from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import text
from app.database import get_db
from pydantic import BaseModel
from typing import Optional
from datetime import date
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

router = APIRouter(prefix="/reports", tags=["reports"])


class GenerateRequest(BaseModel):
    date_from: date
    date_to: date
    purchaser_ids: Optional[list[int]] = None


def _get_eligible_payouts(db: Session, date_from: date, date_to: date):
    """Return marketplace payout rows eligible for commission calculation.
    Parking passes are included - they're real sold/paid-out inventory and
    go through the same commission flow as any other sale.

    Effective date precedence: the row's own payment date (only some
    marketplaces provide one), then the file-level payment date entered at
    import (same source the ReachPro push and Adjustments views use), and
    only then the import timestamp - otherwise rows from date-less
    marketplaces are invisible to any range that ends before their import
    day."""
    rows = db.execute(text("""
        SELECT
            mp.payout_id,
            mp.sale_id,
            mp.amount,
            mp.payment_type,
            mp.adjustment_commission_treatment,
            mp.marketplace_order_id,
            mp.payment_date,
            mp.created_at,
            s.proceeds,
            s.cost,
            s.pnl,
            s.tags,
            s.purchased_date,
            e.performer,
            e.event_tag,
            COALESCE(mp.payment_date, pif.payment_date, mp.created_at::date) AS effective_date
        FROM marketplace_payouts mp
        JOIN sales s ON s.sale_id = mp.sale_id
        LEFT JOIN events e ON e.event_id = s.event_id
        LEFT JOIN payout_import_files pif ON pif.import_file_id = mp.import_file_id
        WHERE mp.sale_id IS NOT NULL
          AND mp.commission_included_in_payout_id IS NULL
          AND COALESCE(mp.payment_date, pif.payment_date, mp.created_at::date) BETWEEN :d_from AND :d_to
          AND s.pnl IS NOT NULL
    """), {"d_from": date_from, "d_to": date_to}).fetchall()
    return [dict(r._mapping) for r in rows]


def _get_purchasers_for_sale(db: Session, sale_id: int):
    rows = db.execute(text("""
        SELECT sp.purchaser_id, sp.split_pct, p.name, p.default_commission_pct
        FROM sale_purchasers sp
        JOIN purchasers p ON p.purchaser_id = sp.purchaser_id
        WHERE sp.sale_id = :sid
          AND p.default_commission_pct > 0
    """), {"sid": sale_id}).fetchall()
    return [dict(r._mapping) for r in rows]


def _get_overrides(db: Session, purchaser_id: int):
    rows = db.execute(text("""
        SELECT event_tag, purchased_before, purchased_from, commission_pct
        FROM commission_overrides WHERE purchaser_id = :pid
        ORDER BY (event_tag IS NULL), override_id
    """), {"pid": purchaser_id}).fetchall()
    return [
        {"tag": (r.event_tag or "").strip().lower() or None,
         "purchased_before": r.purchased_before,
         "purchased_from": r.purchased_from,
         "pct": float(r.commission_pct)}
        for r in rows
    ]


def _resolve_commission(overrides: list, performer: str, event_tag: str, default_pct: float,
                        sale_tags: str = "", purchased_date=None) -> float:
    """Return the applicable commission rate for a sale.

    A rule matches when ALL of its set conditions hold:
      - tag: case-insensitive substring of performer name + event tag + sale tags
      - purchased_before: sale purchased earlier (a sale with NO purchase date
        counts as infinitely old, so it always satisfies this)
      - purchased_from: sale purchased on/after (never satisfied without a date)
    Tag-bearing rules are checked before date-only rules (_get_overrides
    orders them); first match wins, else the purchaser default."""
    search = " ".join(filter(None, [performer, event_tag, sale_tags])).lower()
    for rule in overrides:
        if rule["tag"] and rule["tag"] not in search:
            continue
        if rule["purchased_before"] is not None:
            if purchased_date is not None and purchased_date >= rule["purchased_before"]:
                continue
        if rule["purchased_from"] is not None:
            if purchased_date is None or purchased_date < rule["purchased_from"]:
                continue
        if not (rule["tag"] or rule["purchased_before"] or rule["purchased_from"]):
            continue
        return rule["pct"]
    return default_pct


ADJUSTMENT_REASON_LABELS = {
    "missed_ihd": "Missed IHD",
    "doe": "DOE",
    "cancelled_event": "Cancelled Event",
    "rejected_replacements": "Rejected Replacements",
    "other": "Other",
}


def _categorize(marketplace: str, payment_type: str) -> str:
    """Classify a payout line as Sales / Chargebacks / Wastage.

    Chargebacks are identified by payment_type='Adjustment' (a marketplace-issued
    clawback), not by the sign of the amount — a normal completed sale can still
    have a negative P&L without being a chargeback.
    """
    if marketplace == "Wastage":
        return "Wastage"
    if payment_type == "Adjustment":
        return "Chargebacks"
    return "Sales"


def _adjustment_pnl(amount: float, cost: float, treatment: str | None) -> float:
    """PnL contribution of an Adjustment row for the selected commission
    treatment. Unset (None) defaults to 'full_amount', preserving the original
    behavior for adjustments that haven't been reviewed yet."""
    treatment = treatment or "full_amount"
    if treatment == "none":
        return 0.0
    if treatment == "full_amount_minus_cost":
        return amount - cost if amount >= 0 else amount + cost
    return amount


def _pnl_for_row(payment_type: str, amount: float, cost: float, adjustment_treatment: str | None) -> float:
    """Unified PnL-contribution calculation for a single payout row."""
    if payment_type == "Adjustment":
        return _adjustment_pnl(amount, cost, adjustment_treatment)
    if payment_type == "Payment":
        return amount - cost
    return amount  # Wastage and anything else: full amount, no cost netting


def _build_preview(db: Session, date_from: date, date_to: date) -> dict:
    payouts = _get_eligible_payouts(db, date_from, date_to)

    # Cache overrides per purchaser
    override_cache = {}

    # purchaser_id → { name, lines: [...], gross, commission }
    summary: dict = {}

    for p in payouts:
        sale_purchasers = _get_purchasers_for_sale(db, p["sale_id"])
        for sp in sale_purchasers:
            pid = sp["purchaser_id"]
            if pid not in override_cache:
                override_cache[pid] = _get_overrides(db, pid)
            overrides = override_cache[pid]

            commission_pct = _resolve_commission(
                overrides,
                p["performer"] or "",
                p["event_tag"] or "",
                float(sp["default_commission_pct"]),
                p.get("tags") or "",
                p.get("purchased_date"),
            )

            payout_share = float(p["amount"]) * float(sp["split_pct"])
            pnl_per_unit = _pnl_for_row(
                p["payment_type"], float(p["amount"]), float(p["cost"]), p.get("adjustment_commission_treatment")
            )
            pnl_share = pnl_per_unit * float(sp["split_pct"])
            commission_amount = round(pnl_share * commission_pct, 2)

            if pid not in summary:
                summary[pid] = {
                    "purchaser_id": pid,
                    "purchaser_name": sp["name"],
                    "gross_payout": 0.0,
                    "gross_pnl": 0.0,
                    "commission_amount": 0.0,
                    "payout_count": 0,
                    "lines": [],
                    # Expenses aren't tied to an event, so their commission
                    # impact uses the purchaser's default rate (overrides
                    # are event-scoped and can't apply)
                    "default_commission_pct": float(sp["default_commission_pct"]),
                }

            summary[pid]["gross_payout"] = round(summary[pid]["gross_payout"] + payout_share, 2)
            summary[pid]["gross_pnl"] = round(summary[pid]["gross_pnl"] + pnl_share, 2)
            summary[pid]["commission_amount"] = round(summary[pid]["commission_amount"] + commission_amount, 2)
            summary[pid]["payout_count"] += 1
            summary[pid]["lines"].append({
                "payout_id": p["payout_id"],
                "sale_id": p["sale_id"],
                "payout_amount": payout_share,
                "commission_pct_applied": commission_pct,
                "commission_amount": commission_amount,
            })

    # Attach each purchaser's unclaimed expenses. Anything dated on or before
    # the window's end is swept in (not just the window itself) so an expense
    # entered late can never be orphaned between payout runs. Claimed
    # expenses (already in a live batch) are excluded - rollback unclaims.
    #
    # Expenses reduce P&L, not commission dollar-for-dollar: each expense's
    # commission impact is amount x the purchaser's default commission pct.
    for pid, s in summary.items():
        pct = s["default_commission_pct"]
        exp_rows = db.execute(text("""
            SELECT expense_id, expense_date, description, amount
            FROM purchaser_expenses
            WHERE purchaser_id = :pid AND included_in_payout_id IS NULL
              AND expense_date <= :date_to
            ORDER BY expense_date, expense_id
        """), {"pid": pid, "date_to": date_to}).fetchall()
        s["expenses"] = [
            {"expense_id": e.expense_id, "expense_date": str(e.expense_date),
             "description": e.description, "amount": float(e.amount),
             "commission_impact": round(float(e.amount) * pct, 2)}
            for e in exp_rows
        ]
        s["expenses_total"] = round(sum(float(e.amount) for e in exp_rows), 2)
        s["expense_commission_impact"] = round(sum(e["commission_impact"] for e in s["expenses"]), 2)
        s["net_pnl"] = round(s["gross_pnl"] - s["expenses_total"], 2)
        s["net_commission"] = round(s["commission_amount"] - s["expense_commission_impact"], 2)

    # Offline (private) sales whose proceeds are still a $0/$1 placeholder in
    # ReachPro. Under standard treatment they'd compute as a total loss and
    # claw back commission, so the UI warns before generation: fix the price
    # in ReachPro, re-run Sync, and the sweep re-prices the row automatically.
    unpriced_offline = [dict(r._mapping) for r in db.execute(text("""
        SELECT s.reachpro_sale_id, s.sale_date, s.cost, mp.amount AS proceeds,
               e.performer AS event_name,
               STRING_AGG(p.name, ' + ' ORDER BY p.name) AS purchasers
        FROM marketplace_payouts mp
        JOIN marketplaces m ON m.marketplace_id = mp.marketplace_id
        JOIN sales s ON s.sale_id = mp.sale_id
        LEFT JOIN events e ON e.event_id = s.event_id
        JOIN sale_purchasers sp ON sp.sale_id = s.sale_id
        JOIN purchasers p ON p.purchaser_id = sp.purchaser_id
        WHERE m.name = 'Offline' AND mp.amount <= 1 AND s.cost > 0
          AND COALESCE(mp.payment_date, mp.created_at::date) BETWEEN :date_from AND :date_to
        GROUP BY s.reachpro_sale_id, s.sale_date, s.cost, mp.amount, e.performer
        ORDER BY s.sale_date
    """), {"date_from": date_from, "date_to": date_to}).fetchall()]

    # Payments on CANCELLED sales that have no offsetting adjustment yet.
    # Automatiq/B2B files carry the payout but never the clawback, so without
    # a manually-added offset line the purchaser earns commission on a sale
    # that delivered nothing. Surfaced here so a human verifies the offset
    # exists (or is genuinely not coming) before commission locks.
    cancelled_unoffset = [dict(r._mapping) for r in db.execute(text("""
        SELECT s.reachpro_sale_id, s.sale_date, s.fulfillment_status,
               MIN(mp.marketplace_order_id) AS order_id,
               e.performer AS event_name,
               SUM(mp.amount) AS net_unoffset,
               STRING_AGG(DISTINCT p.name, ' + ') AS purchasers
        FROM marketplace_payouts mp
        JOIN sales s ON s.sale_id = mp.sale_id
        LEFT JOIN events e ON e.event_id = s.event_id
        JOIN sale_purchasers sp ON sp.sale_id = s.sale_id
        JOIN purchasers p ON p.purchaser_id = sp.purchaser_id
        LEFT JOIN payout_import_files pif ON pif.import_file_id = mp.import_file_id
        WHERE (s.fulfillment_status ILIKE '%cancel%' OR s.cancellation_date IS NOT NULL)
          AND s.offset_resolution IS NULL
          AND s.sale_id IN (
              SELECT mp2.sale_id FROM marketplace_payouts mp2
              LEFT JOIN payout_import_files pif2 ON pif2.import_file_id = mp2.import_file_id
              WHERE mp2.commission_included_in_payout_id IS NULL
                AND COALESCE(mp2.payment_date, pif2.payment_date, mp2.created_at::date)
                    BETWEEN :date_from AND :date_to
          )
        GROUP BY s.reachpro_sale_id, s.sale_date, s.fulfillment_status, e.performer
        HAVING SUM(mp.amount) > 0.01
        ORDER BY SUM(mp.amount) DESC
    """), {"date_from": date_from, "date_to": date_to}).fetchall()]

    return {
        "date_from": str(date_from),
        "date_to": str(date_to),
        "total_payout_rows": len(payouts),
        "purchasers": list(summary.values()),
        "unpriced_offline": unpriced_offline,
        "cancelled_unoffset": cancelled_unoffset,
    }


@router.get("/preview")
def preview(date_from: date, date_to: date, db: Session = Depends(get_db)):
    return _build_preview(db, date_from, date_to)


@router.get("/purchaser-detail")
def purchaser_detail(purchaser_id: int, date_from: date, date_to: date, db: Session = Depends(get_db)):
    """Category breakdown (Sales / Chargebacks / Wastage) for a single purchaser in a date range."""
    rows = db.execute(text("""
        SELECT
            m.name AS marketplace,
            mp.payment_type,
            mp.adjustment_commission_treatment,
            mp.amount * sp.split_pct AS payout_share,
            s.cost * sp.split_pct AS cost_share,
            s.tags,
            s.purchased_date,
            e.performer,
            e.event_tag
        FROM marketplace_payouts mp
        JOIN sales s ON s.sale_id = mp.sale_id
        JOIN sale_purchasers sp ON sp.sale_id = s.sale_id AND sp.purchaser_id = :pid
        JOIN marketplaces m ON m.marketplace_id = mp.marketplace_id
        JOIN purchasers p ON p.purchaser_id = sp.purchaser_id
        LEFT JOIN events e ON e.event_id = s.event_id
        LEFT JOIN payout_import_files pif ON pif.import_file_id = mp.import_file_id
        WHERE mp.commission_included_in_payout_id IS NULL
          AND COALESCE(mp.payment_date, pif.payment_date, mp.created_at::date) BETWEEN :d_from AND :d_to
          AND s.pnl IS NOT NULL
          AND p.default_commission_pct > 0
    """), {"pid": purchaser_id, "d_from": date_from, "d_to": date_to}).fetchall()

    purchaser = db.execute(text(
        "SELECT default_commission_pct FROM purchasers WHERE purchaser_id = :pid"
    ), {"pid": purchaser_id}).fetchone()
    overrides = _get_overrides(db, purchaser_id)

    buckets = {
        "Sales":       {"count": 0, "total_amount": 0.0, "total_pnl": 0.0, "commission": 0.0},
        "Chargebacks": {"count": 0, "total_amount": 0.0, "total_pnl": 0.0, "commission": 0.0},
        "Wastage":     {"count": 0, "total_amount": 0.0, "total_pnl": 0.0, "commission": 0.0},
    }

    for r in rows:
        payout = float(r.payout_share or 0)
        cost = float(r.cost_share or 0)
        pnl = _pnl_for_row(r.payment_type, payout, cost, r.adjustment_commission_treatment)
        comm_pct = _resolve_commission(
            overrides, r.performer or "", r.event_tag or "",
            float(purchaser.default_commission_pct),
            r.tags or "",
            r.purchased_date,
        )

        key = _categorize(r.marketplace, r.payment_type)

        buckets[key]["count"] += 1
        buckets[key]["total_amount"] += payout
        buckets[key]["total_pnl"] += pnl
        buckets[key]["commission"] += pnl * comm_pct

    return [
        {
            "category": key,
            "count": b["count"],
            "total_amount": round(b["total_amount"], 2),
            "total_pnl": round(b["total_pnl"], 2),
            "commission": round(b["commission"], 2),
        }
        for key, b in buckets.items()
        if b["count"] > 0
    ]


@router.post("/generate")
def generate(body: GenerateRequest, db: Session = Depends(get_db)):
    preview = _build_preview(db, body.date_from, body.date_to)
    # Filter to only requested purchasers if specified
    if body.purchaser_ids:
        preview["purchasers"] = [p for p in preview["purchasers"] if p["purchaser_id"] in body.purchaser_ids]

    if not preview["purchasers"]:
        raise HTTPException(status_code=400, detail="No eligible payout rows found for the selected date range.")

    created_payouts = []

    for p in preview["purchasers"]:
        # Create purchaser_payout header
        pp_id = db.execute(text("""
            INSERT INTO purchaser_payouts (
                purchaser_id, status, payout_date,
                gross_pnl, commission_pct, commission_amount,
                total_other_deductions, expense_commission_impact
            ) VALUES (
                :pid, 'pending_review', :pdate,
                :gross_pnl, NULL, :comm, :expenses, :impact
            ) RETURNING purchaser_payout_id
        """), {
            "pid": p["purchaser_id"],
            "pdate": body.date_to,
            "gross_pnl": p["gross_pnl"],
            "comm": p["commission_amount"],
            "expenses": p.get("expenses_total", 0),
            "impact": p.get("expense_commission_impact", 0),
        }).fetchone()[0]

        # Claim this batch's expenses - same locking discipline as payout
        # lines. Rollback releases them; while claimed they're read-only.
        expense_ids = [e["expense_id"] for e in p.get("expenses", [])]
        if expense_ids:
            db.execute(text("""
                UPDATE purchaser_expenses SET included_in_payout_id = :ppid
                WHERE expense_id = ANY(:eids) AND included_in_payout_id IS NULL
            """), {"ppid": pp_id, "eids": expense_ids})

        # Insert lines
        for line in p["lines"]:
            db.execute(text("""
                INSERT INTO purchaser_payout_lines (
                    purchaser_payout_id, marketplace_payout_id, sale_id,
                    payout_amount, commission_pct_applied, commission_amount
                ) VALUES (
                    :ppid, :mpid, :sid, :amount, :pct, :comm
                ) ON CONFLICT DO NOTHING
            """), {
                "ppid": pp_id,
                "mpid": line["payout_id"],
                "sid": line["sale_id"],
                "amount": line["payout_amount"],
                "pct": line["commission_pct_applied"],
                "comm": line["commission_amount"],
            })

            # Stamp the marketplace_payout to prevent double-paying
            db.execute(text("""
                UPDATE marketplace_payouts
                SET commission_included_in_payout_id = :ppid
                WHERE payout_id = :mpid
            """), {"ppid": pp_id, "mpid": line["payout_id"]})

        created_payouts.append({
            "purchaser_payout_id": pp_id,
            "purchaser_name": p["purchaser_name"],
            "gross_payout": p["gross_payout"],
            "commission_amount": p["commission_amount"],
            "expenses_total": p.get("expenses_total", 0),
            "expense_commission_impact": p.get("expense_commission_impact", 0),
            "net_commission": p.get("net_commission", p["commission_amount"]),
            "payout_count": p["payout_count"],
        })

    db.commit()
    return {"created": created_payouts}


def _get_remaining_inventory(db: Session, buyer_user_id) -> list:
    """Read this purchaser's future-event inventory from the local cache table,
    refreshed via POST /sync/inventory-refresh (see sync.py)."""
    rows = db.execute(text("""
        SELECT event_name, qty, cost, last_purchase_date
        FROM remaining_inventory
        WHERE buyer_user_id = :bid
        ORDER BY event_name
    """), {"bid": str(buyer_user_id)}).fetchall()
    return [dict(r._mapping) for r in rows]


@router.get("/purchaser-payouts/{payout_id}/download")
def download_purchaser_payout(payout_id: int, db: Session = Depends(get_db)):
    # Header
    header = db.execute(text("""
        SELECT pp.purchaser_payout_id, pp.payout_date, pp.commission_amount,
               pp.expense_commission_impact,
               p.name AS purchaser_name, p.buyer_user_id
        FROM purchaser_payouts pp
        JOIN purchasers p ON p.purchaser_id = pp.purchaser_id
        WHERE pp.purchaser_payout_id = :id
    """), {"id": payout_id}).fetchone()
    if not header:
        raise HTTPException(status_code=404, detail="Payout not found")

    # Detailed lines
    lines = db.execute(text("""
        SELECT
            mk.name AS marketplace,
            mp.payment_type,
            mp.adjustment_commission_treatment,
            mp.adjustment_reason,
            mp.marketplace_order_id AS order_id,
            ppl.sale_id,
            ppl.payout_amount AS amount_paid_out,
            ppl.commission_pct_applied,
            ppl.commission_amount,
            e.performer AS event_name,
            e.event_date,
            e.event_time,
            e.venue,
            s.quantity AS qty,
            s.proceeds AS sale_price,
            s.cost,
            s.pnl,
            s.sale_date,
            s.purchased_date,
            pif.filename AS source_file,
            sp.split_pct
        FROM purchaser_payout_lines ppl
        JOIN marketplace_payouts mp ON mp.payout_id = ppl.marketplace_payout_id
        JOIN marketplaces mk ON mk.marketplace_id = mp.marketplace_id
        JOIN sales s ON s.sale_id = ppl.sale_id
        LEFT JOIN events e ON e.event_id = s.event_id
        LEFT JOIN payout_import_files pif ON pif.import_file_id = mp.import_file_id
        JOIN sale_purchasers sp ON sp.sale_id = s.sale_id
            AND sp.purchaser_id = (
                SELECT purchaser_id FROM purchaser_payouts WHERE purchaser_payout_id = :id
            )
        WHERE ppl.purchaser_payout_id = :id
    """), {"id": payout_id}).fetchall()

    # Expenses claimed by this batch - itemized on the Summary tab and
    # subtracted from total commission
    expenses = [dict(r._mapping) for r in db.execute(text("""
        SELECT expense_date, description, amount
        FROM purchaser_expenses
        WHERE included_in_payout_id = :id
        ORDER BY expense_date, expense_id
    """), {"id": payout_id}).fetchall()]

    inventory = _get_remaining_inventory(db, header.buyer_user_id) if header.buyer_user_id else []
    # synced_at on the cache rows is written only on a fully successful
    # refresh, so MAX of it is "when this inventory data was last updated"
    inventory_synced_at = db.execute(text(
        "SELECT MAX(synced_at) FROM remaining_inventory"
    )).scalar() if inventory else None

    # Adjustments treated as "none" are excluded from commission entirely and
    # are hidden from the purchaser's export - they still exist in the DB and
    # remain visible in the Review Adjustments page.
    line_dicts = [
        dict(r._mapping) for r in lines
        if not (r.payment_type == "Adjustment" and r.adjustment_commission_treatment == "none")
    ]

    wb = _build_workbook(line_dicts, inventory, inventory_synced_at, expenses,
                         float(header.expense_commission_impact or 0))

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    safe_name = header.purchaser_name.replace(" ", "_")
    filename = f"{safe_name}_commission_{header.payout_date}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _strip_tz(dt):
    """Remove timezone info so openpyxl can write the datetime."""
    if dt is None:
        return None
    if hasattr(dt, 'tzinfo') and dt.tzinfo is not None:
        return dt.replace(tzinfo=None)
    return dt


def _build_workbook(lines: list, inventory: list = None, inventory_synced_at=None,
                    expenses: list = None, expense_impact: float = 0.0) -> openpyxl.Workbook:
    wb = openpyxl.Workbook()

    # ── Raw Data sheet ──────────────────────────────────────────────────────
    ws_raw = wb.active
    ws_raw.title = "Raw Data"

    raw_headers = [
        "Event Name", "Event Date", "Venue", "Order ID", "Payout Type", "Adjustment Reason",
        "Payout Amount", "Qty", "Sale Price", "Cost", "P&L", "Commission", "Sale Date",
        "Purchase Date", "Source File",
    ]
    _write_header_row(ws_raw, raw_headers)

    highlight_fill = PatternFill("solid", fgColor="FCE4E4")

    from datetime import datetime as dt_cls

    def _sort_key(r):
        category = _categorize(r.get("marketplace"), r.get("payment_type"))
        return (
            r.get("event_name") or "",
            r.get("event_date") or date.min,
            r.get("order_id") or "",
            category,
        )

    for r in sorted(lines, key=_sort_key):
        event_dt = None
        if r.get("event_date"):
            ed = r["event_date"]
            et = r.get("event_time")
            event_dt = dt_cls.combine(ed, et) if et else ed

        split = float(r.get("split_pct") or 1)
        sale_price = float(r.get("sale_price") or 0) * split
        cost = float(r.get("cost") or 0) * split
        amount_paid_out = float(r.get("amount_paid_out") or 0)
        pnl = _pnl_for_row(r.get("payment_type"), amount_paid_out, cost, r.get("adjustment_commission_treatment"))
        commission = float(r.get("commission_amount") or 0)
        category = _categorize(r.get("marketplace"), r.get("payment_type"))
        reason_label = (
            ADJUSTMENT_REASON_LABELS.get(r.get("adjustment_reason"), "")
            if r.get("payment_type") == "Adjustment" else ""
        )

        ws_raw.append([
            r.get("event_name"),
            event_dt,
            r.get("venue"),
            r.get("order_id"),
            category,
            reason_label,
            amount_paid_out,
            r.get("qty"),
            float(sale_price),
            float(cost),
            pnl,
            commission,
            _strip_tz(r.get("sale_date")),
            r.get("purchased_date"),
            r.get("source_file"),
        ])

        if category in ("Chargebacks", "Wastage"):
            row_idx = ws_raw.max_row
            for col in range(1, len(raw_headers) + 1):
                ws_raw.cell(row=row_idx, column=col).fill = highlight_fill

    # Totals row
    last = ws_raw.max_row
    total_row = last + 1
    bold = Font(bold=True)
    ws_raw.cell(row=total_row, column=1, value="TOTAL").font = bold
    for col, formula_col in [(7, "G"), (9, "I"), (10, "J"), (11, "K"), (12, "L")]:
        cell = ws_raw.cell(row=total_row, column=col,
                           value=f"=SUM({formula_col}2:{formula_col}{last})")
        cell.font = bold
        cell.number_format = '#,##0.00'

    _auto_col_widths(ws_raw)
    _format_currency_cols(ws_raw, [7, 9, 10, 11, 12], start_row=2)

    # ── Summary sheet ───────────────────────────────────────────────────────
    ws_sum = wb.create_sheet("Summary")

    # -- Category breakdown (Sales / Chargebacks / Wastage) --
    cat_headers = ["Category", "# Rows", "Total Amount", "Total P&L", "Commission"]
    _write_header_row(ws_sum, cat_headers, row=1)

    cat_buckets = {
        "Sales": {"count": 0, "total_amount": 0.0, "total_pnl": 0.0, "commission": 0.0},
        "Chargebacks": {"count": 0, "total_amount": 0.0, "total_pnl": 0.0, "commission": 0.0},
        "Wastage": {"count": 0, "total_amount": 0.0, "total_pnl": 0.0, "commission": 0.0},
    }
    for r in lines:
        split = float(r.get("split_pct") or 1)
        amount = float(r.get("amount_paid_out") or 0)
        cost = float(r.get("cost") or 0) * split
        pnl = _pnl_for_row(r.get("payment_type"), amount, cost, r.get("adjustment_commission_treatment"))
        b = cat_buckets[_categorize(r.get("marketplace"), r.get("payment_type"))]
        b["count"] += 1
        b["total_amount"] += amount
        b["total_pnl"] += pnl
        b["commission"] += float(r.get("commission_amount") or 0)

    for key, b in cat_buckets.items():
        if b["count"] > 0:
            ws_sum.append([key, b["count"], round(b["total_amount"], 2), round(b["total_pnl"], 2), round(b["commission"], 2)])

    # Expenses sit in the same table as ordinary rows: the expense dollars hit
    # the P&L column (negative) and the commission column carries the
    # commission-rate impact of that amount, so the TOTAL row nets both
    # automatically. Per-row impact allocates the batch's frozen total impact
    # proportionally by amount (uniform rate at generation makes this exact),
    # with any rounding drift folded into the last row so the column sums to
    # exactly the stored total.
    if expenses:
        amounts = [float(e["amount"]) for e in expenses]
        total_amt = sum(amounts)
        impacts = [
            round(expense_impact * a / total_amt, 2) if total_amt else 0.0
            for a in amounts
        ]
        if impacts:
            impacts[-1] = round(impacts[-1] + (round(expense_impact, 2) - round(sum(impacts), 2)), 2)
        for e, impact in zip(expenses, impacts):
            ws_sum.append([
                f"Expense ({e['expense_date']}): {e['description']}",
                None, None,
                -float(e["amount"]),
                -impact,
            ])

    cat_last = ws_sum.max_row
    cat_total_row = cat_last + 1
    bold = Font(bold=True)
    ws_sum.cell(row=cat_total_row, column=1, value="TOTAL").font = bold
    for col, letter in [(3, "C"), (4, "D"), (5, "E")]:
        cell = ws_sum.cell(row=cat_total_row, column=col,
                           value=f"=SUM({letter}2:{letter}{cat_last})")
        cell.font = bold
        cell.number_format = '#,##0.00'
    _format_currency_cols(ws_sum, [3, 4, 5], start_row=2)

    # -- Event breakdown --
    event_header_row = cat_total_row + 2
    sum_headers = [
        "Event Name", "Venue", "Event Date", "# Tickets",
        "Sale Price", "Amount Paid Out", "Cost", "P&L", "Commission",
    ]
    _write_header_row(ws_sum, sum_headers, row=event_header_row)

    from collections import defaultdict
    events: dict = defaultdict(lambda: {
        "tickets": 0, "sale_price": 0.0, "amount_paid_out": 0.0,
        "cost": 0.0, "pnl": 0.0, "commission": 0.0,
        "_seen_sales": set(),
    })
    for r in lines:
        # Grouped by event + venue + date, not event name alone - two matches
        # sharing a generic name (e.g. multiple "World Cup" games at different
        # stadiums/dates) must stay on separate rows, not collapse into one.
        key = (r.get("event_name") or "(Unknown)", r.get("venue"), r.get("event_date"))
        ev = events[key]
        split = float(r.get("split_pct") or 1)
        sale_id = r.get("sale_id")

        # Amount paid out, P&L and commission accumulate per payout line - see
        # _pnl_for_row for how cost nets against each payment_type, so a sale
        # with several rows (e.g. an Adjustment alongside the original Payment)
        # doesn't have its cost subtracted more than once.
        amount_line = float(r.get("amount_paid_out") or 0)
        row_cost = float(r.get("cost") or 0) * split
        ev["amount_paid_out"] += amount_line
        ev["pnl"] += _pnl_for_row(r.get("payment_type"), amount_line, row_cost, r.get("adjustment_commission_treatment"))
        ev["commission"] += float(r.get("commission_amount") or 0)

        # Cost (display only), sale_price and qty are per-sale — only count once per sale_id
        if sale_id not in ev["_seen_sales"]:
            ev["_seen_sales"].add(sale_id)
            ev["tickets"] += int(r.get("qty") or 0)
            ev["sale_price"] += float(r.get("sale_price") or 0) * split
            ev["cost"] += float(r.get("cost") or 0) * split

    sorted_events = sorted(events.items(), key=lambda x: x[1]["sale_price"], reverse=True)

    for (event_name, venue, event_date), ev in sorted_events:
        pnl = round(ev["pnl"], 2)
        ws_sum.append([
            event_name,
            venue,
            event_date,
            ev["tickets"],
            round(ev["sale_price"], 2),
            round(ev["amount_paid_out"], 2),
            round(ev["cost"], 2),
            pnl,
            round(ev["commission"], 2),
        ])

    # Totals row
    last_sum = ws_sum.max_row
    total_sum = last_sum + 1
    event_data_start = event_header_row + 1
    ws_sum.cell(row=total_sum, column=1, value="TOTAL").font = bold
    for col, letter in [(4, "D"), (5, "E"), (6, "F"), (7, "G"), (8, "H"), (9, "I")]:
        cell = ws_sum.cell(row=total_sum, column=col,
                           value=f"=SUM({letter}{event_data_start}:{letter}{last_sum})")
        cell.font = bold
        if col != 4:
            cell.number_format = '#,##0.00'

    _auto_col_widths(ws_sum)
    _format_currency_cols(ws_sum, [5, 6, 7, 8, 9], start_row=event_data_start)

    # ── Remaining Inventory sheet ────────────────────────────────────────────
    if inventory:
        ws_inv = wb.create_sheet("Remaining Inventory")

        stamp = ""
        if inventory_synced_at is not None:
            stamp = f"Inventory last updated: {inventory_synced_at.strftime('%Y-%m-%d %H:%M')} UTC"
        stamp_cell = ws_inv.cell(row=1, column=1, value=stamp or "Remaining Inventory")
        stamp_cell.font = Font(italic=True, color="666666")

        inv_headers = ["Event Name", "Qty", "Cost", "Last Purchased"]
        _write_header_row(ws_inv, inv_headers, row=2)

        for ev in inventory:
            ws_inv.append([
                ev["event_name"],
                ev["qty"],
                round(ev["cost"], 2),
                ev.get("last_purchase_date"),
            ])

        inv_last = ws_inv.max_row
        inv_total_row = inv_last + 1
        ws_inv.cell(row=inv_total_row, column=1, value="TOTAL").font = bold
        for col, letter in [(2, "B"), (3, "C")]:
            cell = ws_inv.cell(row=inv_total_row, column=col,
                               value=f"=SUM({letter}3:{letter}{inv_last})")
            cell.font = bold
            if col == 3:
                cell.number_format = '#,##0.00'

        _auto_col_widths(ws_inv)
        _format_currency_cols(ws_inv, [3], start_row=3)

    # Move Summary to first position
    wb.move_sheet("Summary", offset=-1)

    return wb


def _write_header_row(ws, headers, row=1):
    header_fill = PatternFill("solid", fgColor="4472C4")
    header_font = Font(bold=True, color="FFFFFF")
    header_align = Alignment(horizontal="center", wrap_text=True)
    thin = Side(style="thin")
    border = Border(bottom=thin)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = border

    ws.row_dimensions[row].height = 30


def _auto_col_widths(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                lines = str(cell.value or "").split("\n")
                max_len = max(max_len, max(len(l) for l in lines))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 45)


def _format_currency_cols(ws, col_indices, start_row=2):
    fmt = '#,##0.00'
    for col in col_indices:
        for row in range(start_row, ws.max_row + 1):
            ws.cell(row=row, column=col).number_format = fmt


@router.get("/purchaser-payouts")
def list_purchaser_payouts(include_rolled_back: bool = False, db: Session = Depends(get_db)):
    query = """
        SELECT
            pp.purchaser_payout_id,
            p.name AS purchaser_name,
            pp.status,
            pp.payout_date,
            pp.gross_pnl,
            pp.commission_amount,
            pp.total_other_deductions AS expenses_total,
            pp.expense_commission_impact,
            pp.commission_amount - pp.expense_commission_impact AS net_commission,
            pp.created_at,
            COUNT(ppl.id) AS line_count
        FROM purchaser_payouts pp
        JOIN purchasers p ON p.purchaser_id = pp.purchaser_id
        LEFT JOIN purchaser_payout_lines ppl ON ppl.purchaser_payout_id = pp.purchaser_payout_id
    """
    if not include_rolled_back:
        query += " WHERE pp.status != 'rolled_back'"
    query += """
        GROUP BY pp.purchaser_payout_id, p.name
        ORDER BY pp.created_at DESC
    """
    rows = db.execute(text(query)).fetchall()
    return [dict(r._mapping) for r in rows]


@router.post("/purchaser-payouts/{payout_id}/commit")
def commit_payout(payout_id: int, db: Session = Depends(get_db)):
    pp = db.execute(text(
        "SELECT status FROM purchaser_payouts WHERE purchaser_payout_id = :id"
    ), {"id": payout_id}).fetchone()
    if not pp:
        raise HTTPException(status_code=404, detail="Payout not found")
    if pp.status == "committed":
        raise HTTPException(status_code=400, detail="Already committed")
    if pp.status == "rolled_back":
        raise HTTPException(status_code=400, detail="Cannot commit a rolled-back payout")
    db.execute(text(
        "UPDATE purchaser_payouts SET status = 'committed' WHERE purchaser_payout_id = :id"
    ), {"id": payout_id})
    db.commit()
    return {"status": "committed"}


@router.post("/purchaser-payouts/{payout_id}/rollback")
def rollback_payout(payout_id: int, db: Session = Depends(get_db)):
    pp = db.execute(text(
        "SELECT status FROM purchaser_payouts WHERE purchaser_payout_id = :id"
    ), {"id": payout_id}).fetchone()
    if not pp:
        raise HTTPException(status_code=404, detail="Payout not found")
    if pp.status == "committed":
        raise HTTPException(status_code=400, detail="Cannot roll back a committed payout")
    if pp.status == "rolled_back":
        raise HTTPException(status_code=400, detail="Already rolled back")

    # Unlock the marketplace payout rows
    db.execute(text("""
        UPDATE marketplace_payouts
        SET commission_included_in_payout_id = NULL
        WHERE commission_included_in_payout_id = :id
    """), {"id": payout_id})

    # Release the batch's expenses back to unclaimed so they're editable
    # again and get swept into the next generation
    db.execute(text("""
        UPDATE purchaser_expenses
        SET included_in_payout_id = NULL
        WHERE included_in_payout_id = :id
    """), {"id": payout_id})

    # Delete the lines
    db.execute(text(
        "DELETE FROM purchaser_payout_lines WHERE purchaser_payout_id = :id"
    ), {"id": payout_id})

    # Mark header rolled back
    db.execute(text(
        "UPDATE purchaser_payouts SET status = 'rolled_back' WHERE purchaser_payout_id = :id"
    ), {"id": payout_id})

    db.commit()
    return {"status": "rolled_back"}
